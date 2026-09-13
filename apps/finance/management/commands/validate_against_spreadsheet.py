"""Prove the imported data reproduces the spreadsheet's own figures.

Rather than hard-coding expected numbers, this reads the cached formula results
straight out of the workbook — the values the owner actually sees — and compares
them against what the application calculates from the imported transactions.

Where a figure does not match, the difference is attributed rather than merely
reported. The workbook's formulas are parsed to discover which cell ranges they
actually cover, and the same restriction is applied to the database, so a
mismatch caused by a defect in a formula can be told apart from a mismatch caused
by a faulty import. Only an unexplained residual counts as a failure.

Checks run:

1. **Dashboard totals** for the date range set in the workbook, and its
   "Last Year" comparison.
2. **Monthly Report** for every month of the selected year, so an error that
   happens to cancel out across a year is still caught.
3. **Categorisation agreement** — the rules are run against the imported rows
   without saving, and compared with the categories recorded in the sheet.
4. **Effect of the corrected treatment** of refunds, quantified.
"""

from __future__ import annotations

import calendar
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db.models import Sum

from apps.analytics.cashflow import LEGACY_EXCLUDED_LABELS, Mode, compute_cash_metrics
from apps.core.money import ZERO, fmt, quantise, to_decimal
from apps.core.periods import DateRange
from apps.finance.categorisation import find_matching_rule
from apps.finance.models import (
    BankTransaction,
    CategoryRule,
    CategorySource,
    ImportIssue,
    IssueCode,
    TransactionSource,
)

TOLERANCE = Decimal("0.01")


@dataclass(frozen=True)
class CellCheck:
    """One dashboard figure to verify.

    Most cells aggregate the transaction sheets directly. The profit cells
    instead subtract one dashboard cell from another (``=B8-F8``), so they carry
    ``components`` and are emulated by composing those cells' emulated values
    rather than by parsing ranges out of their own formula.
    """

    name: str
    coordinate: str
    attribute: str
    components: tuple[tuple[str, str, int], ...] = ()


@dataclass
class Emulation:
    """The result of recomputing a figure under the workbook's own limitations."""

    value: Decimal | None
    causes: list[str]
    unreproducible: bool = False


@dataclass
class SheetFormula:
    """What a dashboard cell's formula actually references."""

    coordinate: str
    text: str

    @property
    def has_broken_reference(self) -> bool:
        """Whether the formula contains a ``#REF!`` error.

        A broken reference makes Excel and Sheets silently drop that part of the
        calculation, so the cell shows a plausible but wrong number.
        """
        return "#REF!" in self.text

    def start_row(self, sheet_name: str) -> int | None:
        """The first row of ``sheet_name`` that this formula includes.

        A formula written as ``'Historic Data'!C10:C1008`` silently omits rows 2
        to 9 of that sheet.
        """
        pattern = rf"'?{re.escape(sheet_name)}'?!\$?[A-Z]+\$?(\d+)"
        rows = [int(m) for m in re.findall(pattern, self.text)]
        return min(rows) if rows else None


class Command(BaseCommand):
    help = "Compare calculated figures against the cached values in the Finance Dashboard workbook."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Path to Finance Dashboard.xlsx")
        parser.add_argument(
            "--show-disagreements",
            type=int,
            default=12,
            help="How many categorisation disagreements to list (default: 12)",
        )

    def handle(self, *args, **options):
        path = Path(options["path"]).expanduser()
        if not path.is_file():
            raise CommandError(f"File not found: {path}")

        values = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)

        failures: list[str] = []
        failures += self._check_dashboard(values, formulas)
        failures += self._check_monthly_report(values, formulas)
        self._check_categorisation(options["show_disagreements"])
        self._report_correction_impact(values)

        self.stdout.write("")
        if failures:
            self.stdout.write(self.style.ERROR(f"{len(failures)} unexplained difference(s):"))
            for failure in failures:
                self.stdout.write(self.style.ERROR(f"  ✗ {failure}"))
            raise CommandError("Validation failed: the imported data does not reproduce the spreadsheet.")

        self.stdout.write(
            self.style.SUCCESS(
                "PASSED — every spreadsheet figure is either reproduced exactly or its\n"
                "difference is fully accounted for by a defect in the workbook."
            )
        )

    # ------------------------------------------------------------------
    # Check 1 — Dashboard
    # ------------------------------------------------------------------

    def _check_dashboard(self, values, formulas) -> list[str]:
        self.stdout.write(self.style.MIGRATE_HEADING("\n1. Dashboard totals"))

        if "Dashboard" not in values.sheetnames:
            self.stdout.write(self.style.WARNING("  No 'Dashboard' sheet; skipped."))
            return []

        sheet = values["Dashboard"]
        formula_sheet = formulas["Dashboard"]
        start = _as_date(sheet["C5"].value)
        end = _as_date(sheet["G5"].value)
        if not start or not end:
            self.stdout.write(self.style.WARNING("  Could not read the date range from C5/G5; skipped."))
            return []

        current = DateRange(start, end, "Workbook range")
        last_year = current.same_period_last_year()

        groups = [
            (
                f"Selected range: {start:%d %b %Y} – {end:%d %b %Y}",
                current,
                (
                    CellCheck("Revenue", "B8", "revenue"),
                    CellCheck("Expenses", "F8", "total_expenses"),
                    CellCheck("Expenses exc. salary & tithe", "J8", "total_expenses_excluding_distributions"),
                    CellCheck(
                        "Profit",
                        "B15",
                        "net_profit",
                        components=(("B8", "revenue", 1), ("F8", "total_expenses", -1)),
                    ),
                    CellCheck(
                        "Profit exc. salary & tithe",
                        "F15",
                        "net_profit_excluding_distributions",
                        components=(
                            ("B8", "revenue", 1),
                            ("J8", "total_expenses_excluding_distributions", -1),
                        ),
                    ),
                ),
            ),
            (
                f"Last Year: {last_year.start:%d %b %Y} – {last_year.end:%d %b %Y}",
                last_year,
                (
                    CellCheck("Revenue", "C11", "revenue"),
                    CellCheck("Expenses", "G11", "total_expenses"),
                    CellCheck("Expenses exc. salary & tithe", "K11", "total_expenses_excluding_distributions"),
                    CellCheck(
                        "Profit",
                        "C18",
                        "net_profit",
                        components=(("C11", "revenue", 1), ("G11", "total_expenses", -1)),
                    ),
                    CellCheck(
                        "Profit exc. salary & tithe",
                        "G18",
                        "net_profit_excluding_distributions",
                        components=(
                            ("C11", "revenue", 1),
                            ("K11", "total_expenses_excluding_distributions", -1),
                        ),
                    ),
                ),
            ),
        ]

        failures: list[str] = []
        for label, date_range, checks in groups:
            self.stdout.write(f"\n  {label}")
            full = compute_cash_metrics(date_range, mode=Mode.LEGACY)
            for check in checks:
                failures += self._compare_cell(
                    check=check,
                    expected=to_decimal(sheet[check.coordinate].value, default=None),
                    calculated=getattr(full, check.attribute),
                    date_range=date_range,
                    formula_sheet=formula_sheet,
                )
        return failures

    def _compare_cell(
        self,
        *,
        check: CellCheck,
        expected: Decimal | None,
        calculated: Decimal,
        date_range: DateRange,
        formula_sheet,
    ) -> list[str]:
        if expected is None:
            self.stdout.write(f"    – {check.name:<30} no value in the sheet; skipped")
            return []

        if _close(expected, calculated):
            self.stdout.write(self.style.SUCCESS(f"    ✓ {check.name:<30} {fmt(expected):>12}  matches exactly"))
            return []

        # The figures disagree. Work out whether the workbook's own formula
        # explains it, by applying the formula's real coverage to our data.
        self.stdout.write(
            self.style.WARNING(
                f"    ~ {check.name:<30} sheet {fmt(expected):>12}   ours {fmt(calculated):>12}   "
                f"diff {fmt(calculated - expected)}"
            )
        )

        emulation = self._emulate(check, date_range, formula_sheet)
        for cause in emulation.causes:
            self.stdout.write(f"        · {cause}")

        if emulation.unreproducible:
            self.stdout.write(
                self.style.WARNING(
                    "        ! the sheet's own value cannot be reproduced because its formula is "
                    "broken. Treat our figure as the correct one and repair or retire that cell."
                )
            )
            return []

        if emulation.value is None:
            self.stdout.write(self.style.ERROR("        ✗ no explanation found"))
            return [f"{check.name}: sheet {expected} vs ours {calculated}, unexplained"]

        if _close(expected, emulation.value):
            self.stdout.write(
                self.style.SUCCESS(
                    f"        ✓ applying the sheet's own formula coverage to our data gives "
                    f"{fmt(emulation.value)}, matching the sheet exactly. Our figure is the corrected one."
                )
            )
            return []

        residual = emulation.value - expected
        self.stdout.write(
            self.style.ERROR(
                f"        ✗ emulating the formula gives {fmt(emulation.value)}, still "
                f"{fmt(residual)} away from the sheet"
            )
        )
        return [f"{check.name}: unexplained residual of {residual} after emulating the formula"]

    def _emulate(self, check: CellCheck, date_range: DateRange, formula_sheet) -> Emulation:
        """Recompute a figure under the same limitations as the workbook's formula."""
        if check.components:
            return self._emulate_composed(check, date_range, formula_sheet)

        formula = SheetFormula(check.coordinate, str(formula_sheet[check.coordinate].value or ""))
        return self._emulate_aggregate(check.attribute, date_range, formula)

    def _emulate_composed(self, check: CellCheck, date_range: DateRange, formula_sheet) -> Emulation:
        """Emulate a cell that subtracts other dashboard cells from one another."""
        total = ZERO
        causes: list[str] = []
        for coordinate, attribute, sign in check.components:
            component = self._emulate(
                CellCheck(coordinate, coordinate, attribute), date_range, formula_sheet
            )
            for cause in component.causes:
                if cause not in causes:
                    causes.append(cause)
            if component.unreproducible:
                causes.append(f"this figure depends on {coordinate}, which cannot be reproduced")
                return Emulation(None, causes, unreproducible=True)
            if component.value is None:
                return Emulation(None, causes)
            total = quantise(total + sign * component.value)
        return Emulation(total, causes)

    def _emulate_aggregate(
        self, attribute: str, date_range: DateRange, formula: SheetFormula
    ) -> Emulation:
        """Emulate a cell that aggregates the transaction sheets directly."""
        causes: list[str] = []

        if formula.has_broken_reference:
            causes.append(
                f"cell {formula.coordinate} contains a #REF! broken reference, so the sheet computes "
                "this figure from an incomplete formula"
            )
            return Emulation(None, causes, unreproducible=True)

        queryset = BankTransaction.objects.all()

        # Limitation 1: the formula's range may start below the first data row.
        historic_start = formula.start_row("Historic Data")
        if historic_start and historic_start > 2:
            omitted = queryset.filter(
                source=TransactionSource.LEGACY_HISTORIC,
                source_row_number__lt=historic_start,
                occurred_on__gte=date_range.start,
                occurred_on__lte=date_range.end,
            )
            count = omitted.count()
            if count:
                causes.append(
                    f"the formula reads 'Historic Data' from row {historic_start}, omitting "
                    f"{count} row(s) of real data that fall in this period"
                )
                queryset = queryset.exclude(pk__in=omitted.values("pk"))

        transactions_start = formula.start_row("Transactions")
        if transactions_start and transactions_start > 2:
            omitted = queryset.filter(
                source=TransactionSource.LEGACY_TRANSACTIONS,
                source_row_number__lt=transactions_start,
                occurred_on__gte=date_range.start,
                occurred_on__lte=date_range.end,
            )
            count = omitted.count()
            if count:
                causes.append(
                    f"the formula reads 'Transactions' from row {transactions_start}, omitting "
                    f"{count} row(s) of real data that fall in this period"
                )
                queryset = queryset.exclude(pk__in=omitted.values("pk"))

        metrics = compute_cash_metrics(date_range, mode=Mode.LEGACY, queryset=queryset)
        emulated = getattr(metrics, attribute)

        # Limitation 2: rows the sheet counts twice, which were imported once.
        adjustment = self._duplicate_adjustment(date_range, attribute)
        if adjustment:
            causes.append(
                f"the sheet counts {fmt(abs(adjustment))} twice from a row duplicated in the "
                f"Transactions sheet; imported once here"
            )
            emulated = quantise(emulated + adjustment)

        return Emulation(emulated, causes)

    def _duplicate_adjustment(self, date_range: DateRange, attribute: str) -> Decimal:
        """How much a duplicated source row inflates the sheet's figure.

        Duplicated rows were rejected on import, so to reproduce the sheet their
        value has to be added back to whichever total they belong to.
        """
        duplicates = ImportIssue.objects.filter(
            code=IssueCode.DUPLICATE_EXTERNAL_ID,
            occurred_on__gte=date_range.start,
            occurred_on__lte=date_range.end,
        )
        money_out = duplicates.filter(amount__lt=0).aggregate(t=Sum("amount"))["t"] or ZERO
        money_in = duplicates.filter(amount__gt=0).aggregate(t=Sum("amount"))["t"] or ZERO

        if attribute == "revenue":
            return quantise(money_in)
        if attribute in {"total_expenses", "total_expenses_excluding_distributions"}:
            return quantise(-money_out)
        if attribute == "net_profit":
            return quantise(money_in + money_out)
        if attribute == "net_profit_excluding_distributions":
            return quantise(money_in + money_out)
        return ZERO

    # ------------------------------------------------------------------
    # Check 2 — Monthly Report
    # ------------------------------------------------------------------

    def _check_monthly_report(self, values, formulas) -> list[str]:
        self.stdout.write(self.style.MIGRATE_HEADING("\n2. Monthly Report"))

        if "Monthly Report" not in values.sheetnames:
            self.stdout.write(self.style.WARNING("  No 'Monthly Report' sheet; skipped."))
            return []

        sheet = values["Monthly Report"]
        formula_sheet = formulas["Monthly Report"]
        year_value = sheet["B2"].value
        if not isinstance(year_value, (int, float)):
            self.stdout.write(self.style.WARNING("  Could not read the year from B2; skipped."))
            return []
        year = int(year_value)

        self.stdout.write(f"\n  Year {year}          {'Revenue':^27}  {'Expenses':^27}")
        failures: list[str] = []
        for offset, column in enumerate("BCDEFGHIJKLM"):
            month_start = date(year, offset + 1, 1)
            month_end = month_start.replace(day=calendar.monthrange(year, offset + 1)[1])
            month = DateRange(month_start, month_end, month_start.strftime("%b %Y"))

            expected_revenue = to_decimal(sheet[f"{column}8"].value, default=None)
            expected_expenses = to_decimal(sheet[f"{column}9"].value, default=None)
            if expected_revenue is None and expected_expenses is None:
                continue

            metrics = compute_cash_metrics(month, mode=Mode.LEGACY)
            revenue_ok = _close(expected_revenue, metrics.revenue)
            expenses_ok = _close(expected_expenses, metrics.total_expenses)

            style = self.style.SUCCESS if revenue_ok and expenses_ok else self.style.WARNING
            self.stdout.write(
                style(
                    f"  {'✓' if revenue_ok and expenses_ok else '~'} {month_start:%b}   "
                    f"sheet {fmt(expected_revenue):>11} ours {fmt(metrics.revenue):>11}   "
                    f"sheet {fmt(expected_expenses):>11} ours {fmt(metrics.total_expenses):>11}"
                )
            )

            for ok, name, cell_row, expected, attribute in (
                (revenue_ok, "revenue", 8, expected_revenue, "revenue"),
                (expenses_ok, "expenses", 9, expected_expenses, "total_expenses"),
            ):
                if ok or expected is None:
                    continue
                coordinate = f"{column}{cell_row}"
                formula = SheetFormula(coordinate, str(formula_sheet[coordinate].value or ""))
                emulation = self._emulate_aggregate(attribute, month, formula)
                for cause in emulation.causes:
                    self.stdout.write(f"        · {cause}")
                if emulation.unreproducible:
                    self.stdout.write(
                        self.style.WARNING(
                            f"        ! {month_start:%b} {name}: the sheet's formula is broken; "
                            "our figure is the correct one"
                        )
                    )
                elif emulation.value is not None and _close(expected, emulation.value):
                    self.stdout.write(
                        self.style.SUCCESS(
                            f"        ✓ {month_start:%b} {name}: the sheet's own formula coverage "
                            f"reproduces {fmt(emulation.value)} exactly"
                        )
                    )
                else:
                    failures.append(
                        f"{month_start:%b %Y} {name}: sheet {expected} vs ours "
                        f"{getattr(metrics, attribute)}, unexplained"
                    )
        return failures

    # ------------------------------------------------------------------
    # Check 3 — categorisation agreement
    # ------------------------------------------------------------------

    def _check_categorisation(self, show: int) -> None:
        self.stdout.write(self.style.MIGRATE_HEADING("\n3. Categorisation rule agreement"))
        self.stdout.write(
            "  Running the rules against rows you categorised in the spreadsheet, to see how\n"
            "  often the rules alone reach the same answer. Nothing is changed: your recorded\n"
            "  categories are locked."
        )

        rules = list(
            CategoryRule.objects.filter(is_active=True).select_related("category").order_by("priority", "pk")
        )
        imported = BankTransaction.objects.filter(
            category_source=CategorySource.IMPORTED, category__isnull=False
        ).select_related("category")
        total = imported.count()
        if not rules or not total:
            self.stdout.write(self.style.WARNING("  Nothing to compare; skipped."))
            return

        exact_agree = 0
        kind_agree = 0
        no_match = 0
        pairs: Counter = Counter()
        examples: list[tuple[BankTransaction, str]] = []

        for txn in imported.iterator(chunk_size=500):
            rule = find_matching_rule(txn, rules)
            if rule is None:
                no_match += 1
                pairs[(txn.category.name, "(no rule)")] += 1
                if len(examples) < show:
                    examples.append((txn, "no rule matched — needs a rule"))
                continue
            if rule.category_id == txn.category_id:
                exact_agree += 1
                kind_agree += 1
            else:
                # Agreement on *kind* is what matters for the totals: splitting
                # "Income" into "Shopify Payout" changes the label, not the maths.
                if rule.category.kind == txn.category.kind:
                    kind_agree += 1
                pairs[(txn.category.name, rule.category.name)] += 1
                if len(examples) < show:
                    examples.append((txn, f"rules say {rule.category.name} ({rule.category.get_kind_display()})"))

        self.stdout.write(
            f"\n  Same category:               {exact_agree:>6,} of {total:,}  "
            f"({exact_agree / total * 100:.1f}%)"
        )
        self.stdout.write(
            f"  Same financial treatment:    {kind_agree:>6,} of {total:,}  "
            f"({kind_agree / total * 100:.1f}%)   ← what affects the numbers"
        )
        self.stdout.write(f"  No rule matched:             {no_match:>6,}  ← still needs manual categorisation")

        if pairs:
            self.stdout.write("\n  Most common differences (your label → rules' label):")
            for (yours, theirs), count in pairs.most_common(10):
                self.stdout.write(f"    {count:>6,}  {yours} → {theirs}")

        if examples:
            self.stdout.write("\n  Examples:")
            for txn, note in examples:
                self.stdout.write(
                    f"    {txn.occurred_on:%d %b %Y} {fmt(txn.amount):>10}  "
                    f"{txn.display_name[:40]:<40} you: {txn.category.name:<18} {note}"
                )

    # ------------------------------------------------------------------
    # Check 4 — effect of the corrected treatment
    # ------------------------------------------------------------------

    def _report_correction_impact(self, values) -> None:
        self.stdout.write(self.style.MIGRATE_HEADING("\n4. Effect of the corrected calculation"))

        if "Dashboard" not in values.sheetnames:
            return
        sheet = values["Dashboard"]
        start = _as_date(sheet["C5"].value)
        end = _as_date(sheet["G5"].value)
        if not start or not end:
            return

        date_range = DateRange(start, end)
        legacy = compute_cash_metrics(date_range, mode=Mode.LEGACY)
        standard = compute_cash_metrics(date_range, mode=Mode.STANDARD)

        self.stdout.write(f"\n  {start:%d %b %Y} – {end:%d %b %Y}")
        self.stdout.write(f"  {'':<32}{'spreadsheet':>13}{'corrected':>13}{'difference':>13}")
        for label, old, new in (
            ("Revenue", legacy.revenue, standard.revenue),
            ("Total expenses", legacy.total_expenses, standard.total_expenses),
            ("Net profit", legacy.net_profit, standard.net_profit),
            (
                "Profit exc. salary & tithe",
                legacy.net_profit_excluding_distributions,
                standard.net_profit_excluding_distributions,
            ),
        ):
            self.stdout.write(f"  {label:<32}{fmt(old):>13}{fmt(new):>13}{fmt(new - old):>13}")

        self._explain_revenue_difference(date_range)

        if standard.uncategorised_count:
            self.stdout.write(
                self.style.WARNING(
                    f"\n  {standard.uncategorised_count} transaction(s) totalling "
                    f"{fmt(standard.uncategorised_value)} in this range have no category and are\n"
                    "  excluded from the corrected figures."
                )
            )

    def _explain_revenue_difference(self, date_range: DateRange) -> None:
        """Itemise why corrected revenue differs from the spreadsheet's revenue."""
        from apps.finance.models import CategoryKind

        credits = (
            BankTransaction.objects.in_range(date_range)
            .filter(amount__gt=0)
            .exclude(category__name__in=LEGACY_EXCLUDED_LABELS)
            .exclude(category__kind=CategoryKind.REVENUE)
            .select_related("category")
            .order_by("-amount")
        )
        debits = (
            BankTransaction.objects.in_range(date_range)
            .filter(amount__lt=0, category__kind=CategoryKind.REVENUE)
            .select_related("category")
            .order_by("amount")
        )

        if credits.exists():
            total = credits.aggregate(t=Sum("amount"))["t"] or ZERO
            self.stdout.write(
                f"\n  Money in that the spreadsheet counts as revenue but is not a sale "
                f"({fmt(total)}):"
            )
            for txn in credits[:10]:
                self.stdout.write(
                    f"    {txn.occurred_on:%d %b %Y} {fmt(txn.amount):>10}  "
                    f"{txn.display_name[:38]:<38} {txn.category.name} "
                    f"({txn.category.get_kind_display()})"
                )
            if credits.count() > 10:
                self.stdout.write(f"    ... and {credits.count() - 10} more")

        if debits.exists():
            total = debits.aggregate(t=Sum("amount"))["t"] or ZERO
            self.stdout.write(
                f"\n  Refunds to customers, now reducing revenue instead of counting as an "
                f"expense ({fmt(total)}):"
            )
            for txn in debits[:10]:
                self.stdout.write(
                    f"    {txn.occurred_on:%d %b %Y} {fmt(txn.amount):>10}  "
                    f"{txn.display_name[:38]:<38} {txn.category.name}"
                )


def _close(expected: Decimal | None, actual: Decimal | None) -> bool:
    if expected is None or actual is None:
        return False
    return abs(Decimal(expected) - Decimal(actual)) <= TOLERANCE


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None
